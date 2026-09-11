# -*- coding: utf-8 -*-
"""
03_make_paper_results_CLAUDE_FINAL.py
========================
Research-only pipeline 3/3. No Streamlit/web UI.

Reads scripts 01/02 outputs and creates publication-oriented tables and 300-dpi figures.

Main outputs
------------
paper_output/PAPER_LLM_RULE_CAS_GAP_RESULTS.xlsx
paper_output/tables/Table1_benchmark_composition.csv
paper_output/tables/Table2_threeway_performance.csv
paper_output/tables/Table3_performance_by_scope.csv
paper_output/tables/Table4_hybrid_marginal_benefit.csv
paper_output/tables/Table5_failure_summary.csv
paper_output/tables/Table6_official_API_PDF_audit.csv
paper_output/tables/Table7_LLM_runtime_reproducibility.csv
paper_output/figures/Fig1_study_design.png
paper_output/figures/Fig2_benchmark_scope_composition.png
paper_output/figures/Fig3_threeway_recall.png
paper_output/figures/Fig4_false_safe_rate.png
paper_output/figures/Fig5_scope_recall_heatmap.png
paper_output/figures/Fig6_capability_performance.png
paper_output/figures/Fig7_rule_marginal_benefit_by_capability.png
paper_output/RESULTS_README.txt

Dependencies
------------
pip install pandas numpy openpyxl matplotlib

Run
---
python 03_make_paper_results_CLAUDE_FINAL.py
"""
from __future__ import annotations

import json
import os
from pathlib import Path
from typing import Optional

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parent
INTERMEDIATE = ROOT / "intermediate"
DATA_DIR = ROOT / "data"
OUT = ROOT / "paper_output"
TABLE_DIR = OUT / "tables"
FIG_DIR = OUT / "figures"
for d in (OUT, TABLE_DIR, FIG_DIR):
    d.mkdir(parents=True, exist_ok=True)

XLSX = OUT / "PAPER_LLM_RULE_CAS_GAP_RESULTS.xlsx"


def read_csv(name: str, required: bool = False) -> pd.DataFrame:
    p = INTERMEDIATE / name
    if not p.exists():
        if required:
            raise FileNotFoundError(f"Missing required file: {p}. Run scripts 01 and 02 first.")
        return pd.DataFrame()
    try:
        return pd.read_csv(p, encoding="utf-8-sig", low_memory=False)
    except pd.errors.EmptyDataError:
        return pd.DataFrame()


def read_data_csv(name: str) -> pd.DataFrame:
    p = DATA_DIR / name
    if not p.exists():
        return pd.DataFrame()
    try:
        return pd.read_csv(p, encoding="utf-8-sig", low_memory=False)
    except pd.errors.EmptyDataError:
        return pd.DataFrame()

def aggregate_performance(perf: pd.DataFrame) -> pd.DataFrame:
    if perf.empty:
        return perf
    metrics = [
        "automation_coverage", "positive_case_coverage", "accuracy", "precision",
        "recall", "specificity", "f1", "false_safe_rate", "false_positive_rate",
    ]
    rows = []
    for (system, model), g in perf.groupby(["system", "llm_model"], dropna=False):
        row = {
            "system": system,
            "llm_model": model,
            "llm_backend": str(g["llm_backend"].dropna().iloc[0]) if "llm_backend" in g.columns and g["llm_backend"].notna().any() else "",
            "capability_tier": str(g["capability_tier"].dropna().iloc[0]) if "capability_tier" in g.columns and g["capability_tier"].notna().any() else "",
            "capability_order": float(pd.to_numeric(g.get("capability_order"), errors="coerce").dropna().iloc[0]) if "capability_order" in g.columns and pd.to_numeric(g["capability_order"], errors="coerce").notna().any() else np.nan,
            "parameter_size_b": float(pd.to_numeric(g.get("parameter_size_b"), errors="coerce").dropna().iloc[0]) if "parameter_size_b" in g.columns and pd.to_numeric(g["parameter_size_b"], errors="coerce").notna().any() else np.nan,
            "n_repeats": int(g["repeat"].nunique()) if "repeat" in g.columns else 1,
            "n_total_benchmark": int(pd.to_numeric(g.get("n_total"), errors="coerce").max()) if "n_total" in g.columns else np.nan,
                        "n_common_mean": float(pd.to_numeric(g.get("n_common"), errors="coerce").mean()) if "n_common" in g.columns else np.nan,
            "common_analysis_coverage_mean": float(pd.to_numeric(g.get("common_analysis_coverage"), errors="coerce").mean()) if "common_analysis_coverage" in g.columns else np.nan,
            "system_output_coverage_mean": float(pd.to_numeric(g.get("system_output_coverage"), errors="coerce").mean()) if "system_output_coverage" in g.columns else np.nan,
        }
        for m in metrics:
            if m in g.columns:
                v = pd.to_numeric(g[m], errors="coerce")
                row[f"{m}_mean"] = float(v.mean()) if v.notna().any() else np.nan
                row[f"{m}_std"] = float(v.std(ddof=1)) if v.notna().sum() > 1 else 0.0
        rows.append(row)
    out = pd.DataFrame(rows)
    order = {"RULE_ENGINE_ONLY": 0, "LLM_ONLY": 1, "LLM_PLUS_RULE": 2}
    out["_order"] = out["system"].map(order).fillna(99)
    out = out.sort_values(["_order", "llm_model"]).drop(columns="_order").reset_index(drop=True)
    return out


def aggregate_by_scope(by: pd.DataFrame) -> pd.DataFrame:
    if by.empty:
        return by
    metrics = ["common_coverage", "accuracy", "precision", "recall", "false_safe_rate", "false_positive_rate"]
    rows = []
    for (system, model, scope), g in by.groupby(["system", "llm_model", "reference_scope_type"], dropna=False):
        row = {
            "system": system, "llm_model": model, "reference_scope_type": scope,
            "llm_backend": str(g["llm_backend"].dropna().iloc[0]) if "llm_backend" in g.columns and g["llm_backend"].notna().any() else "",
            "capability_tier": str(g["capability_tier"].dropna().iloc[0]) if "capability_tier" in g.columns and g["capability_tier"].notna().any() else "",
            "capability_order": float(pd.to_numeric(g.get("capability_order"), errors="coerce").dropna().iloc[0]) if "capability_order" in g.columns and pd.to_numeric(g["capability_order"], errors="coerce").notna().any() else np.nan,
            "n_repeats": int(g["repeat"].nunique()) if "repeat" in g.columns else 1,
                        "n_scope_rows": int(pd.to_numeric(g.get("n_total"), errors="coerce").max()) if "n_total" in g.columns else np.nan,
            "n_common_mean": float(pd.to_numeric(g.get("n_common"), errors="coerce").mean()) if "n_common" in g.columns else np.nan,
        }
        for m in metrics:
            if m in g.columns:
                v = pd.to_numeric(g[m], errors="coerce")
                row[f"{m}_mean"] = float(v.mean()) if v.notna().any() else np.nan
                row[f"{m}_std"] = float(v.std(ddof=1)) if v.notna().sum() > 1 else 0.0
        rows.append(row)
    return pd.DataFrame(rows)


def aggregate_marginal(m: pd.DataFrame) -> pd.DataFrame:
    if m.empty:
        return m
    numeric = [c for c in m.columns if c.startswith("delta_")]
    rows = []
    for model, g in m.groupby("llm_model", dropna=False):
        row = {
            "llm_model": model,
            "llm_backend": str(g["llm_backend"].dropna().iloc[0]) if "llm_backend" in g.columns and g["llm_backend"].notna().any() else "",
            "capability_tier": str(g["capability_tier"].dropna().iloc[0]) if "capability_tier" in g.columns and g["capability_tier"].notna().any() else "",
            "capability_order": float(pd.to_numeric(g.get("capability_order"), errors="coerce").dropna().iloc[0]) if "capability_order" in g.columns and pd.to_numeric(g["capability_order"], errors="coerce").notna().any() else np.nan,
            "parameter_size_b": float(pd.to_numeric(g.get("parameter_size_b"), errors="coerce").dropna().iloc[0]) if "parameter_size_b" in g.columns and pd.to_numeric(g["parameter_size_b"], errors="coerce").notna().any() else np.nan,
            "n_repeats": int(g["repeat"].nunique()) if "repeat" in g.columns else 1,
        }
        for c in numeric:
            v = pd.to_numeric(g[c], errors="coerce")
            row[f"{c}_mean"] = float(v.mean()) if v.notna().any() else np.nan
            row[f"{c}_std"] = float(v.std(ddof=1)) if v.notna().sum() > 1 else 0.0
        rows.append(row)
    return pd.DataFrame(rows)


def benchmark_composition(primary: pd.DataFrame, full_summary: pd.DataFrame) -> pd.DataFrame:
    if primary.empty:
        return full_summary
    p = (
        primary.groupby(["reference_scope_type", "reference_requires_extended_identity"], dropna=False)
        .size().reset_index(name="n_primary")
    )
    if full_summary.empty:
        return p
    keep = [c for c in ["reference_scope_type", "reference_requires_extended_identity", "n_full_rows"] if c in full_summary.columns]
    return full_summary[keep].merge(p, on=["reference_scope_type", "reference_requires_extended_identity"], how="outer").fillna(0)


def failure_summary(pred: pd.DataFrame) -> pd.DataFrame:
    if pred.empty:
        return pd.DataFrame()
    x = pred[pred["prediction_available"].astype(str).str.lower().isin(["true", "1", "yes"])].copy()
    if x.empty:
        return pd.DataFrame()
    x["failure_type"] = np.where(
        x["false_safe"].astype(str).str.lower().isin(["true", "1", "yes"]), "FALSE_SAFE",
        np.where(x["false_positive"].astype(str).str.lower().isin(["true", "1", "yes"]), "FALSE_POSITIVE", "CORRECT")
    )
    return (
        x.groupby(["system", "llm_model", "reference_scope_type", "failure_type"], dropna=False)
        .size().reset_index(name="n")
        .sort_values(["system", "llm_model", "failure_type", "n"], ascending=[True, True, True, False])
    )


def runtime_reproducibility(model_status: pd.DataFrame, consistency: pd.DataFrame, llm_results: pd.DataFrame) -> pd.DataFrame:
    rows = []
    models = sorted(set(model_status.get("llm_model", pd.Series(dtype=str)).dropna().astype(str)) |
                    set(llm_results.get("llm_model", pd.Series(dtype=str)).dropna().astype(str)))
    for model in models:
        ms = model_status[model_status.get("llm_model", pd.Series(index=model_status.index, dtype=str)).astype(str).eq(model)] if len(model_status) else pd.DataFrame()
        lr = llm_results[llm_results.get("llm_model", pd.Series(index=llm_results.index, dtype=str)).astype(str).eq(model)] if len(llm_results) else pd.DataFrame()
        cs = consistency[consistency.get("llm_model", pd.Series(index=consistency.index, dtype=str)).astype(str).eq(model)] if len(consistency) else pd.DataFrame()
        if len(lr):
            if "llm_classification_valid" in lr.columns:
                ok = lr["llm_classification_valid"].fillna(False).astype(bool)
            else:
                ok = lr["llm_status"].astype(str).eq("OK") if "llm_status" in lr.columns else pd.Series(False, index=lr.index)
            evidence_ok = lr["llm_evidence_quote_valid"].fillna(False).astype(bool) if "llm_evidence_quote_valid" in lr.columns else pd.Series(False, index=lr.index)
        else:
            ok = pd.Series(dtype=bool)
            evidence_ok = pd.Series(dtype=bool)
        rows.append({
            "llm_model": model,
            "llm_backend": ms.iloc[0].get("llm_backend", lr.iloc[0].get("llm_backend", "") if len(lr) else "") if len(ms) else (lr.iloc[0].get("llm_backend", "") if len(lr) else ""),
            "capability_tier": ms.iloc[0].get("capability_tier", lr.iloc[0].get("capability_tier", "") if len(lr) else "") if len(ms) else (lr.iloc[0].get("capability_tier", "") if len(lr) else ""),
            "capability_order": ms.iloc[0].get("capability_order", lr.iloc[0].get("capability_order", np.nan) if len(lr) else np.nan) if len(ms) else (lr.iloc[0].get("capability_order", np.nan) if len(lr) else np.nan),
            "backend_status": ms.iloc[0].get("backend_status", "") if len(ms) else "",
            "model_available": ms.iloc[0].get("model_available", "") if len(ms) else "",
            "n_llm_outputs": int(len(lr)),
            "validated_output_rate": float(ok.mean()) if len(ok) else np.nan,
            "classification_valid_output_rate": float(ok.mean()) if len(ok) else np.nan,
            "literal_evidence_exact_match_rate_all": float(evidence_ok.mean()) if len(evidence_ok) else np.nan,
            "literal_evidence_exact_match_rate_classification_valid": float(evidence_ok.loc[ok].mean()) if len(evidence_ok) and len(ok) and ok.any() else np.nan,
            "mean_confidence_valid_outputs": float(pd.to_numeric(lr.loc[ok, "llm_confidence"], errors="coerce").mean()) if len(lr) and ok.any() and "llm_confidence" in lr.columns else np.nan,
            "mean_input_tokens": float(pd.to_numeric(lr.get("input_tokens"), errors="coerce").mean()) if len(lr) and "input_tokens" in lr.columns and pd.to_numeric(lr["input_tokens"], errors="coerce").notna().any() else np.nan,
            "mean_output_tokens": float(pd.to_numeric(lr.get("output_tokens"), errors="coerce").mean()) if len(lr) and "output_tokens" in lr.columns and pd.to_numeric(lr["output_tokens"], errors="coerce").notna().any() else np.nan,
            "repeat_decision_consistency": cs.iloc[0].get("repeat_decision_consistency", np.nan) if len(cs) else np.nan,
        })
    return pd.DataFrame(rows)


def aggregate_capability_common(cap: pd.DataFrame) -> pd.DataFrame:
    if cap.empty:
        return cap
    metrics = ["accuracy", "precision", "recall", "specificity", "f1", "false_safe_rate", "false_positive_rate"]
    rows = []
    for (system, model), g in cap.groupby(["system", "llm_model"], dropna=False):
        row = {
            "system": system, "llm_model": model,
            "llm_backend": str(g["llm_backend"].dropna().iloc[0]) if "llm_backend" in g.columns and g["llm_backend"].notna().any() else "",
            "capability_tier": str(g["capability_tier"].dropna().iloc[0]) if "capability_tier" in g.columns and g["capability_tier"].notna().any() else "",
            "capability_order": float(pd.to_numeric(g["capability_order"], errors="coerce").dropna().iloc[0]) if "capability_order" in g.columns and pd.to_numeric(g["capability_order"], errors="coerce").notna().any() else np.nan,
            "parameter_size_b": float(pd.to_numeric(g["parameter_size_b"], errors="coerce").dropna().iloc[0]) if "parameter_size_b" in g.columns and pd.to_numeric(g["parameter_size_b"], errors="coerce").notna().any() else np.nan,
            "n_cross_model_common": int(pd.to_numeric(g.get("n_cross_model_common"), errors="coerce").max()) if "n_cross_model_common" in g.columns else np.nan,
            "n_repeats": int(g["repeat"].nunique()) if "repeat" in g.columns else 1,
        }
        for m in metrics:
            if m in g.columns:
                v = pd.to_numeric(g[m], errors="coerce")
                row[f"{m}_mean"] = float(v.mean()) if v.notna().any() else np.nan
                row[f"{m}_std"] = float(v.std(ddof=1)) if v.notna().sum() > 1 else 0.0
        rows.append(row)
    return pd.DataFrame(rows).sort_values(["capability_order", "llm_model", "system"], kind="stable")


def aggregate_capability_benefit(x: pd.DataFrame) -> pd.DataFrame:
    if x.empty:
        return x
    numeric = [c for c in x.columns if c.startswith("delta_") or c in {"llm_recall", "rule_recall", "hybrid_recall"}]
    rows = []
    for model, g in x.groupby("llm_model", dropna=False):
        row = {
            "llm_model": model,
            "llm_backend": str(g["llm_backend"].dropna().iloc[0]) if "llm_backend" in g.columns and g["llm_backend"].notna().any() else "",
            "capability_tier": str(g["capability_tier"].dropna().iloc[0]) if "capability_tier" in g.columns and g["capability_tier"].notna().any() else "",
            "capability_order": float(pd.to_numeric(g["capability_order"], errors="coerce").dropna().iloc[0]) if "capability_order" in g.columns and pd.to_numeric(g["capability_order"], errors="coerce").notna().any() else np.nan,
            "parameter_size_b": float(pd.to_numeric(g["parameter_size_b"], errors="coerce").dropna().iloc[0]) if "parameter_size_b" in g.columns and pd.to_numeric(g["parameter_size_b"], errors="coerce").notna().any() else np.nan,
            "n_cross_model_common": int(pd.to_numeric(g.get("n_cross_model_common"), errors="coerce").max()) if "n_cross_model_common" in g.columns else np.nan,
            "n_repeats": int(g["repeat"].nunique()) if "repeat" in g.columns else 1,
        }
        for c in numeric:
            v = pd.to_numeric(g[c], errors="coerce")
            row[f"{c}_mean"] = float(v.mean()) if v.notna().any() else np.nan
            row[f"{c}_std"] = float(v.std(ddof=1)) if v.notna().sum() > 1 else 0.0
        rows.append(row)
    return pd.DataFrame(rows).sort_values(["capability_order", "llm_model"], kind="stable")

def study_notes() -> pd.DataFrame:
    return pd.DataFrame([
        {"item": "study_type", "value": "Controlled benchmark of whether regulatory scope is safely representable by direct-CAS equality alone; research-only, no web UI."},
        {"item": "primary_endpoint", "value": "Detection of regulatory scope that cannot be safely represented by direct-CAS equality alone."},
        {"item": "systems", "value": "RULE_ENGINE_ONLY (direct-CAS presence baseline) vs LLM_ONLY vs LLM_PLUS_RULE."},
        {"item": "refined_reference_rule", "value": "A specific salt/polymer/derivative/reaction product/structural-range substance with its own direct CAS is NOT a positive gap unless the source explicitly extends beyond that CAS."},
        {"item": "positive_examples", "value": "No direct CAS; explicit 'A and its salts/compounds/derivatives'; generic designation plus specific-item list; explicit exclusion clause."},
        {"item": "LLM_role", "value": "Semantic scope interpretation from supplied regulatory source text; not final legal decision-making."},
        {"item": "common_analysis_set", "value": "Main Rule-vs-LLM-vs-Hybrid metrics use exactly the same cases: rows with validated LLM output. Coverage is reported separately."},
        {"item": "rule_full_benchmark", "value": "Rule-only performance on all primary benchmark rows is supplementary and must not be compared numerically with common-set LLM metrics without noting the different analysis set."},
        {"item": "reference_boundary", "value": "Source-derived predefined operational reference; not independent expert gold."},
        {"item": "official_source_boundary", "value": "Law API/PDF extractions are audit/candidate evidence and never auto-promoted to the approved master."},
        {"item": "primary_safety_metric", "value": "False-safe rate among CAS-only-insufficient cases, jointly interpreted with recall and LLM output coverage."},
        {"item": "capability_question", "value": "How does the marginal benefit of deterministic rules change as evaluated LLM capability increases? The study does not claim that LLMs are permanently inferior or superior."},
        {"item": "capability_design", "value": "LOW (Qwen 0.8B), MID (Qwen 4B), HIGH (current strong cloud LLM: Claude Sonnet 5). Capability tiers are ordinal study labels, not a universal ranking."},
        {"item": "cross_model_common_set", "value": "Capability comparisons use only benchmark cases with validated output from every evaluated model/repeat stratum."},
    ])

def setup_plot():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({
        "font.size": 10,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "figure.dpi": 120,
        "savefig.bbox": "tight",
    })
    return plt


def fig1_study_design(plt) -> Path:
    fig, ax = plt.subplots(figsize=(12, 4.2))
    ax.axis("off")
    boxes = [
        (0.02, 0.34, 0.17, 0.34, "Official sources\nLaw API\nAppendix PDF\nApproved master"),
        (0.26, 0.34, 0.18, 0.34, "Fixed benchmark\nCAS-only sufficient vs\nCAS-only insufficient\nscope"),
        (0.53, 0.61, 0.18, 0.20, "Rule engine only"),
        (0.53, 0.36, 0.18, 0.20, "LLM only"),
        (0.53, 0.11, 0.18, 0.20, "LLM + rule"),
        (0.80, 0.34, 0.18, 0.34, "Evaluation\nRecall / precision\nFalse-safe / coverage\nScope-specific errors"),
    ]
    for x, y, w, h, text in boxes:
        ax.add_patch(plt.Rectangle((x, y), w, h, fill=False, linewidth=1.4, transform=ax.transAxes))
        ax.text(x+w/2, y+h/2, text, ha="center", va="center", transform=ax.transAxes, fontsize=9.5)
    arrows = [((0.19, 0.51), (0.26, 0.51)), ((0.44, 0.51), (0.53, 0.71)), ((0.44, 0.51), (0.53, 0.46)), ((0.44, 0.51), (0.53, 0.21)), ((0.71, 0.71), (0.80, 0.58)), ((0.71, 0.46), (0.80, 0.51)), ((0.71, 0.21), (0.80, 0.44))]
    for a, b in arrows:
        ax.annotate("", xy=b, xytext=a, xycoords=ax.transAxes, textcoords=ax.transAxes, arrowprops=dict(arrowstyle="->", lw=1.2))
    ax.set_title("Study design: CAS-only screening insufficiency benchmark", fontsize=12)
    p = FIG_DIR / "Fig1_study_design.png"
    fig.savefig(p, dpi=300)
    plt.close(fig)
    return p


def fig2_scope_composition(plt, comp: pd.DataFrame) -> Optional[Path]:
    if comp.empty or "reference_scope_type" not in comp.columns:
        return None
    pdat = comp.copy()
    valcol = "n_primary" if "n_primary" in pdat.columns else "n_full_rows"
    pdat[valcol] = pd.to_numeric(pdat[valcol], errors="coerce").fillna(0)
    pdat = pdat.groupby("reference_scope_type", as_index=False)[valcol].sum().sort_values(valcol)
    fig, ax = plt.subplots(figsize=(8.5, max(4.2, 0.45 * len(pdat))))
    ax.barh(np.arange(len(pdat)), pdat[valcol].values)
    ax.set_yticks(np.arange(len(pdat)))
    ax.set_yticklabels(pdat["reference_scope_type"].astype(str))
    ax.set_xlabel("Number of benchmark rows")
    ax.set_title("Primary benchmark composition by regulatory identity scope")
    for i, v in enumerate(pdat[valcol].values):
        ax.text(v, i, f" {int(v)}", va="center", fontsize=9)
    p = FIG_DIR / "Fig2_benchmark_scope_composition.png"
    fig.savefig(p, dpi=300)
    plt.close(fig)
    return p


def _performance_labels(perf: pd.DataFrame) -> pd.Series:
    return np.where(
        perf["system"].eq("RULE_ENGINE_ONLY"),
        "Rule engine",
        perf["system"].str.replace("_", " ") + "\n" + perf["llm_model"].astype(str),
    )


def fig_metric(plt, perf: pd.DataFrame, metric: str, title: str, ylabel: str, filename: str) -> Optional[Path]:
    if perf.empty or metric not in perf.columns:
        return None
    pdat = perf.copy()
    pdat["label"] = _performance_labels(pdat)
    vals = pd.to_numeric(pdat[metric], errors="coerce").to_numpy()
    fig, ax = plt.subplots(figsize=(max(7.2, 1.35 * len(pdat)), 4.6))
    x = np.arange(len(pdat))
    bars = ax.bar(x, vals)
    ax.set_xticks(x)
    ax.set_xticklabels(pdat["label"], rotation=22, ha="right")
    ax.set_ylim(0, 1.05)
    ax.set_ylabel(ylabel)
    ax.set_title(title)
    for b, v in zip(bars, vals):
        if np.isfinite(v):
            ax.text(b.get_x()+b.get_width()/2, v+0.015, f"{v:.2f}", ha="center", fontsize=9)
    p = FIG_DIR / filename
    fig.savefig(p, dpi=300)
    plt.close(fig)
    return p


def fig5_scope_heatmap(plt, by_scope: pd.DataFrame) -> Optional[Path]:
    if by_scope.empty:
        return None
    b = by_scope.copy()
    b = b[b["system"].isin(["RULE_ENGINE_ONLY", "LLM_ONLY", "LLM_PLUS_RULE"])]
    if b.empty or "recall_mean" not in b.columns:
        return None
    b["series"] = np.where(
        b["system"].eq("RULE_ENGINE_ONLY"), "Rule engine",
        b["system"].str.replace("_", " ") + " | " + b["llm_model"].astype(str)
    )
    piv = b.pivot_table(index="reference_scope_type", columns="series", values="recall_mean", aggfunc="mean")
    if piv.empty:
        return None
    fig, ax = plt.subplots(figsize=(max(7.5, 1.15 * len(piv.columns)), max(4.5, 0.55 * len(piv))))
    im = ax.imshow(piv.values.astype(float), aspect="auto", vmin=0, vmax=1)
    ax.set_xticks(np.arange(len(piv.columns)))
    ax.set_xticklabels(piv.columns, rotation=28, ha="right")
    ax.set_yticks(np.arange(len(piv.index)))
    ax.set_yticklabels(piv.index)
    ax.set_title("CAS-gap recall by identity-scope type")
    for i in range(len(piv.index)):
        for j in range(len(piv.columns)):
            v = piv.iloc[i, j]
            ax.text(j, i, "NA" if pd.isna(v) else f"{v:.2f}", ha="center", va="center", fontsize=8)
    fig.colorbar(im, ax=ax, label="Recall")
    p = FIG_DIR / "Fig5_scope_recall_heatmap.png"
    fig.savefig(p, dpi=300)
    plt.close(fig)
    return p


def fig6_capability_performance(plt, cap_perf: pd.DataFrame) -> Optional[Path]:
    if cap_perf.empty or "recall_mean" not in cap_perf.columns:
        return None
    x = cap_perf.copy()
    x = x[x["system"].isin(["RULE_ENGINE_ONLY", "LLM_ONLY", "LLM_PLUS_RULE"])].copy()
    if x.empty:
        return None
    # One x position per model capability. Rule is shown on each matched stratum to keep the paired baseline explicit.
    models = (x[["llm_model", "capability_tier", "capability_order"]]
              .drop_duplicates().sort_values(["capability_order", "llm_model"], kind="stable"))
    labels = [f"{r.capability_tier}\n{r.llm_model}" for r in models.itertuples()]
    pos = np.arange(len(models))
    fig, ax = plt.subplots(figsize=(max(7.5, 2.0 * len(models)), 4.8))
    for system, label in [("RULE_ENGINE_ONLY", "Rule only"), ("LLM_ONLY", "LLM only"), ("LLM_PLUS_RULE", "LLM + rule")]:
        vals = []
        for m in models["llm_model"]:
            g = x[(x["llm_model"] == m) & (x["system"] == system)]
            vals.append(pd.to_numeric(g["recall_mean"], errors="coerce").mean() if len(g) else np.nan)
        ax.plot(pos, vals, marker="o", label=label)
    ax.set_xticks(pos)
    ax.set_xticklabels(labels)
    ax.set_ylim(0, 1.05)
    ax.set_ylabel("Recall on cross-model common set")
    ax.set_title("LLM capability and CAS-gap detection performance")
    ax.legend()
    p = FIG_DIR / "Fig6_capability_performance.png"
    fig.savefig(p, dpi=300)
    plt.close(fig)
    return p


def fig7_rule_marginal_by_capability(plt, cap_benefit: pd.DataFrame) -> Optional[Path]:
    metric = "delta_rule_benefit_recall_mean"
    if cap_benefit.empty or metric not in cap_benefit.columns:
        return None
    x = cap_benefit.sort_values(["capability_order", "llm_model"], kind="stable").copy()
    vals = pd.to_numeric(x[metric], errors="coerce").to_numpy()
    pos = np.arange(len(x))
    labels = [f"{r.capability_tier}\n{r.llm_model}" for r in x.itertuples()]
    fig, ax = plt.subplots(figsize=(max(7.2, 2.0 * len(x)), 4.6))
    ax.plot(pos, vals, marker="o")
    ax.axhline(0, linewidth=0.8)
    ax.set_xticks(pos)
    ax.set_xticklabels(labels)
    ax.set_ylabel("Δ recall (LLM + rule − LLM only)")
    ax.set_title("Marginal benefit of deterministic rules across LLM capability")
    for i, v in enumerate(vals):
        if np.isfinite(v):
            ax.text(i, v, f"{v:+.3f}", ha="center", va="bottom" if v >= 0 else "top", fontsize=9)
    p = FIG_DIR / "Fig7_rule_marginal_benefit_by_capability.png"
    fig.savefig(p, dpi=300)
    plt.close(fig)
    return p

def autosize_workbook(writer) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col_idx, col in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in list(col)[:1000]:
                val = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(val), 70))
                if len(val) > 40:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 60)
        ws.auto_filter.ref = ws.dimensions


def write_excel(sheets: list[tuple[str, pd.DataFrame]]) -> None:
    with pd.ExcelWriter(XLSX, engine="openpyxl") as w:
        for name, df in sheets:
            (df if df is not None else pd.DataFrame()).to_excel(w, sheet_name=name[:31], index=False)
        autosize_workbook(w)


def main() -> None:
    print("=" * 80)
    print("03 / PAPER TABLES + FIGURES")
    print("=" * 80)

    primary = read_csv("01_cas_gap_benchmark_primary.csv", required=True)
    full_summary = read_csv("01_benchmark_scope_summary.csv")
    trace = read_csv("01_source_traceability_audit.csv")
    api_pdf = read_csv("01_official_api_pdf_status.csv")
    appendix = read_csv("01_official_appendix_pdf_candidates.csv")

    pred = read_csv("02_threeway_predictions.csv", required=True)
    perf_raw = read_csv("02_threeway_performance.csv")
    by_raw = read_csv("02_threeway_by_scope.csv")
    marginal_raw = read_csv("02_hybrid_marginal_benefit.csv")
    paired = read_csv("02_paired_mcnemar_tests.csv")
    consistency = read_csv("02_llm_repeat_consistency.csv")
    model_status = read_csv("02_model_runtime_status.csv")
    llm_results = read_csv("02_llm_semantic_results.csv")
    false_safe = read_csv("02_false_safe_cases.csv")
    false_positive = read_csv("02_false_positive_cases.csv")
    rule_full = read_csv("02_rule_full_benchmark_performance.csv")
    ref_reclass = read_csv("01_reference_reclassification_audit.csv")
    excluded_inactive = read_csv("01_benchmark_excluded_inactive_rows.csv")
    capability_common_raw = read_csv("02_capability_common_performance.csv")
    capability_benefit_raw = read_csv("02_capability_rule_benefit.csv")
    cross_model_common_ids = read_csv("02_cross_model_common_ids.csv")
    master_only_pairs = read_csv("01_master_only_pairs_vs_official_xlsx.csv")
    official_only_pairs = read_csv("01_official_only_pairs_vs_master.csv")
    pair_diff_summary = read_csv("01_official_master_pair_diff_summary.csv")

    t1 = benchmark_composition(primary, full_summary)
    t2 = aggregate_performance(perf_raw)
    t3 = aggregate_by_scope(by_raw)
    t4 = aggregate_marginal(marginal_raw)
    t5 = failure_summary(pred)
    t6 = api_pdf.copy()
    t7 = runtime_reproducibility(model_status, consistency, llm_results)
    t10 = aggregate_capability_common(capability_common_raw)
    t11 = aggregate_capability_benefit(capability_benefit_raw)

    tables = {
        "Table1_benchmark_composition.csv": t1,
        "Table2_threeway_performance.csv": t2,
        "Table3_performance_by_scope.csv": t3,
        "Table4_hybrid_marginal_benefit.csv": t4,
        "Table5_failure_summary.csv": t5,
        "Table6_official_API_PDF_audit.csv": t6,
        "Table7_LLM_runtime_reproducibility.csv": t7,
        "Table8_rule_full_benchmark_supplementary.csv": rule_full,
        "Table9_reference_reclassification_audit.csv": ref_reclass,
        "Table10_capability_common_performance.csv": t10,
        "Table11_capability_rule_benefit.csv": t11,
        "Table12_excluded_inactive_rows.csv": excluded_inactive,
        "Table13_official_master_pair_QC.csv": pair_diff_summary,
        "Table14_master_only_pairs.csv": master_only_pairs,
        "Table15_official_only_pairs.csv": official_only_pairs,
    }
    for name, df in tables.items():
        df.to_csv(TABLE_DIR / name, index=False, encoding="utf-8-sig")

    # Supplementary source-supported broad-salt materials already curated in user's project.
    broad_rules = read_data_csv("broad_salt_rules.csv")
    broad_validation = read_data_csv("broad_salt_validation_cases.csv")

    sheets = [
        ("0_Study_notes", study_notes()),
        ("Table1_Benchmark", t1),
        ("Table2_Performance", t2),
        ("Table3_By_scope", t3),
        ("Table4_Hybrid_gain", t4),
        ("Table5_Failure_summary", t5),
        ("Table6_API_PDF", t6),
        ("Official_Master_pair_QC", pair_diff_summary),
        ("Master_only_pairs_QC", master_only_pairs),
        ("Official_only_pairs_QC", official_only_pairs),
        ("Table7_LLM_status", t7),
        ("Table8_Rule_full_supp", rule_full),
        ("Reference_QC_reclassified", ref_reclass),
        ("Excluded_inactive_rows", excluded_inactive),
        ("Capability_common_perf", t10),
        ("Capability_rule_benefit", t11),
        ("Cross_model_common_IDs", cross_model_common_ids),
        ("Paired_McNemar", paired),
        ("Source_traceability", trace),
        ("Supp_false_safe", false_safe),
        ("Supp_false_positive", false_positive),
        ("Supp_appendix_candidates", appendix),
        ("Supp_primary_benchmark", primary),
        ("Supp_broad_salt_rules", broad_rules),
        ("Supp_broad_salt_validation", broad_validation),
    ]
    write_excel(sheets)

    plt = setup_plot()
    figs = []
    for p in [
        fig1_study_design(plt),
        fig2_scope_composition(plt, t1),
        fig_metric(plt, t2, "recall_mean", "Detection of scope not representable by direct CAS alone", "CAS-only insufficiency recall", "Fig3_threeway_recall.png"),
        fig_metric(plt, t2, "false_safe_rate_mean", "Missed extended-identity cases", "False-safe rate", "Fig4_false_safe_rate.png"),
        fig5_scope_heatmap(plt, t3),
        fig6_capability_performance(plt, t10),
        fig7_rule_marginal_by_capability(plt, t11),
    ]:
        if p is not None:
            figs.append(p)

    manifest = pd.DataFrame([{"figure": p.stem, "file": str(p.relative_to(OUT))} for p in figs])
    manifest.to_csv(OUT / "Figure_manifest.csv", index=False, encoding="utf-8-sig")

    # A concise text summary for checking the run without opening Excel.
    lines = [
        "LLM + RULE ENGINE CAS-GAP STUDY - RESULT OUTPUT",
        "=" * 60,
        f"Primary benchmark rows: {len(primary):,}",
        f"Systems in predictions: {', '.join(sorted(pred['system'].dropna().astype(str).unique())) if 'system' in pred.columns else ''}",
        "",
        "Primary interpretation boundary:",
        "This study evaluates regulatory scope that is not safely representable by direct-CAS equality alone.",
        "The benchmark reference is source-derived and predefined; it is not independent expert gold. Specific complex chemicals with their own CAS are not treated as gaps unless the source explicitly extends beyond that CAS.",
        "Main three-way metrics use the same common validated-LLM cases for Rule, LLM and Hybrid. Law API/PDF candidates remain separate audit evidence.",
        "Cross-capability comparisons additionally use the strict cross-model common set: only cases with valid output from every evaluated model/repeat.",
        "Primary future-proof question: how the marginal recall benefit of deterministic rules changes as evaluated LLM capability increases.",
        "",
        "Main performance table:",
        t2.to_string(index=False) if len(t2) else "No performance rows.",
        "",
        f"Excel: {XLSX}",
        f"Tables: {TABLE_DIR}",
        f"Figures: {FIG_DIR}",
    ]
    (OUT / "RESULTS_README.txt").write_text("\n".join(lines), encoding="utf-8")

    print(f"Excel: {XLSX}")
    print(f"Tables: {TABLE_DIR}")
    print(f"Figures: {FIG_DIR}")
    print(f"Figures generated: {len(figs)}")
    if len(pair_diff_summary):
        q = pair_diff_summary.iloc[0]
        print("\n[Official Excel vs approved master pair QC]")
        for col in [
            "official_unique_pairs", "master_unique_pairs", "pair_overlap",
            "master_only_pairs", "official_only_pairs", "bidirectional_exact_pair_match",
            "official_pair_coverage_by_master", "master_pair_coverage_by_official",
        ]:
            if col in pair_diff_summary.columns:
                print(f"{col}: {q.get(col)}")
    if len(t2):
        show = [c for c in ["system", "llm_model", "n_common_mean", "common_analysis_coverage_mean", "system_output_coverage_mean", "accuracy_mean", "precision_mean", "recall_mean", "false_safe_rate_mean"] if c in t2.columns]
        print("\n[Paper-facing performance]")
        print(t2[show].round(4).to_string(index=False))
    if len(t11):
        print("\n[Capability / rule marginal benefit]")
        show2 = [c for c in ["capability_tier", "llm_model", "n_cross_model_common", "llm_recall_mean", "rule_recall_mean", "hybrid_recall_mean", "delta_rule_benefit_recall_mean"] if c in t11.columns]
        print(t11[show2].round(4).to_string(index=False))


if __name__ == "__main__":
    main()
